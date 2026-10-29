#!/usr/bin/env python3
"""Pull the aircraft table (Directions!AJ1:AU10) and device table (Directions!AA1:AD34) out of the workbook
into site/data/static.json.  Usage: extract_static.py AMR_Template_AZ.xlsm ../site/data/static.json"""
import sys, json, warnings
warnings.filterwarnings('ignore')
import openpyxl
src, out = sys.argv[1], sys.argv[2]
ws = openpyxl.load_workbook(src, data_only=True, keep_vba=True)['Directions']
def v(c):
    x = ws[c].value
    if isinstance(x, float) and x == int(x): x = int(x)
    return x
acft = []
for r in range(2, 11):
    if ws[f'AJ{r}'].value is None: continue
    acft.append(dict(design=v(f'AJ{r}'), mil=v(f'AK{r}'), td=v(f'AL{r}'), gw=v(f'AM{r}'), burn=v(f'AN{r}'), unit=v(f'AO{r}'),
        fob=round(float(ws[f'AP{r}'].value), 2), cr90=v(f'AQ{r}'), cr110=v(f'AR{r}'), cr120=v(f'AS{r}'), hover=v(f'AT{r}'), climb=v(f'AU{r}')))
dev = []
for r in range(2, 35):
    if ws[f'AA{r}'].value is None: continue
    dev.append(dict(device=v(f'AA{r}'), type=v(f'AB{r}'), n=str(v(f'AC{r}')), pad=v(f'AD{r}')))
# DD 1801 code lists + the workbook's default selections ('DD 1801' BB:BD, "XX" marks in BC)
w2 = openpyxl.load_workbook(src, keep_vba=True)['DD 1801']
def codes(a, b):
    out = []
    for r in range(a, b + 1):
        c, d = w2.cell(r, 54).value, w2.cell(r, 56).value
        if c and d:
            c = str(c).strip(); d = str(d).strip()
            out.append([c, d.split('=', 1)[1].strip() if d.startswith(c + '=') else d])
    return out
eq = dict(nav=codes(5, 41), sur=codes(44, 61), pbn=codes(87, 111),
          navDefault=[str(w2.cell(r, 54).value).strip() for r in range(5, 42) if w2.cell(r, 55).value == 'XX'],
          surDefault=[str(w2.cell(r, 54).value).strip() for r in range(44, 62) if w2.cell(r, 55).value == 'XX'],
          pbnDefault=[str(w2.cell(r, 54).value).strip() for r in range(87, 112) if w2.cell(r, 55).value == 'XX'])
# v1.3: only UH-72, UH-60L, UH-60M; devices are plain UH-72 tail numbers (no colour labels). The UH-60L row is the workbook's UH-60A row (no L row exists).
import copy
by = {a['design']: a for a in acft}; L = copy.deepcopy(by['UH-60A']); L['design'] = 'UH-60L'
L['note'] = 'UH-60L: performance row copied from the workbook UH-60A row (no UH-60L row in the workbook) - verify against the -10'
acft = [by['UH-72'], L, by['UH-60M']]
dev = [dict(device=n, type='UH-72', n=n) for n in sorted({d['n'] for d in dev if d['type'] == 'UH-72'})]
json.dump(dict(acft=acft, devices=dev, eq=eq), open(out, 'w'), indent=1)
print(len(acft), 'aircraft,', len(dev), 'devices,', len(eq['nav']), 'nav codes')
